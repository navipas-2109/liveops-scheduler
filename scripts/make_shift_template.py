"""Builds public/templates/shift-import-template.xlsx, the starting file for
the monthly Producer shift sheet that Staff Shift > "นำเข้ากะ" reads.

Run after changing shift codes (keep in step with src/domain/shift-sheet.ts):
    python3 scripts/make_shift_template.py

The layout must stay readable by parseShiftSheet: a header row with
"Employee ID", "Name", "Title", "Squad" followed by the day numbers 1..31.
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "public" / "templates" / "shift-import-template.xlsx"

# code, meaning, hours, fill colour (light, readable on paper and screen)
CODES = [
    ("1", "วันหยุด", "—", "D9D9D9"),
    ("2", "กะเช้า", "07:30–16:30", "CFE2F3"),
    ("4", "กะกลาง", "11:30–20:30", "D9EAD3"),
    ("3", "กะเย็น", "16:30–00:30", "D9D2E9"),
    ("OF", "เวลางานออฟฟิศ (ไม่ได้คุมไลฟ์)", "09:00–18:00", "FCE5CD"),
    ("VL", "ลาพักร้อน", "—", "F4CCCC"),
]
POSITIONS = ["Producer Team Lead", "Senior Producer", "Producer", "Jr. Producer"]

HEADER_ROW = 4
FIRST_DATA_ROW = 5
DATA_ROWS = 200
FIRST_DAY_COL = 5  # E
DAYS = 31

NAVY = "1C2333"
BLUE = "1268FF"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Arial"


def font(**kw):
    return Font(name=FONT, **kw)


def day_col(d):
    return get_column_letter(FIRST_DAY_COL + d - 1)


def date_formula(d):
    return f"DATE(YEAR($B$2),MONTH($B$2),{d})"


def build_shift_sheet(ws, month, rows=()):
    """The grid itself. `rows` pre-fills people (used for the example tab)."""
    last_col = FIRST_DAY_COL + DAYS - 1
    last_row = FIRST_DATA_ROW + DATA_ROWS - 1

    ws["A1"] = "ตารางกะ Producer — LiveOps Scheduler"
    ws["A1"].font = font(bold=True, size=14, color=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

    ws["A2"] = "เดือน"
    ws["A2"].font = font(bold=True)
    ws["A2"].alignment = Alignment(horizontal="right")
    ws["B2"] = month
    ws["B2"].number_format = "mmmm yyyy"
    ws["B2"].font = font(bold=True, color=BLUE, size=12)
    ws["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["B2"].border = BOX
    ws["C2"] = "← ใส่วันที่ 1 ของเดือน เช่น 1/11/2026 (วันในแถว 3 จะเปลี่ยนตาม)"
    ws["C2"].font = font(italic=True, color="7F7F7F", size=9)

    ws["D3"] = "วัน"
    ws["D3"].font = font(bold=True, color="7F7F7F")
    ws["D3"].alignment = Alignment(horizontal="right")

    headers = ["Employee ID", "Name", "Title", "Squad"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font = font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX

    for d in range(1, DAYS + 1):
        col = FIRST_DAY_COL + d - 1
        # Weekday (blank for days the month doesn't have).
        w = ws.cell(3, col, f'=IF(MONTH({date_formula(d)})=MONTH($B$2),CHOOSE(WEEKDAY({date_formula(d)}),"อา","จ","อ","พ","พฤ","ศ","ส"),"")')
        w.font = font(size=9, color="595959")
        w.alignment = Alignment(horizontal="center")
        h = ws.cell(HEADER_ROW, col, d)  # must stay a plain number for the import
        h.font = font(bold=True, color=WHITE)
        h.fill = PatternFill("solid", fgColor=NAVY)
        h.alignment = Alignment(horizontal="center")
        h.border = BOX
        ws.column_dimensions[get_column_letter(col)].width = 4.6

    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 19
    ws.column_dimensions["D"].width = 26
    ws.row_dimensions[HEADER_ROW].height = 20

    for r in range(FIRST_DATA_ROW, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(r, c)
            cell.border = BOX
            cell.font = font(size=10)
            if c >= FIRST_DAY_COL:
                cell.alignment = Alignment(horizontal="center")
        ws.cell(r, 1).number_format = "@"  # keep IDs as typed (e.g. leading zeros)

    for r, (eid, name, title, squad, codes) in enumerate(rows, start=FIRST_DATA_ROW):
        ws.cell(r, 1, eid)
        ws.cell(r, 2, name)
        ws.cell(r, 3, title)
        ws.cell(r, 4, squad or None)
        for d, code in enumerate(codes, start=1):
            if code:
                ws.cell(r, FIRST_DAY_COL + d - 1, int(code) if code.isdigit() else code)

    grid = f"{day_col(1)}{FIRST_DATA_ROW}:{day_col(DAYS)}{last_row}"
    codes_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(c for c, *_ in CODES) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="รหัสกะไม่ถูกต้อง",
        error="ใช้ได้เฉพาะ 1, 2, 3, 4, OF, VL หรือเว้นว่าง (ยังไม่ลงกะ)",
        promptTitle="รหัสกะ",
        prompt="1 หยุด · 2 เช้า · 4 กลาง · 3 เย็น · OF ออฟฟิศ · VL ลา",
        showInputMessage=True,
    )
    ws.add_data_validation(codes_dv)
    codes_dv.add(grid)

    title_dv = DataValidation(
        type="list",
        formula1='"' + ",".join(POSITIONS) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="ตำแหน่งไม่ถูกต้อง",
        error="เลือกจากรายการ: " + ", ".join(POSITIONS),
    )
    ws.add_data_validation(title_dv)
    title_dv.add(f"C{FIRST_DATA_ROW}:C{last_row}")

    # Colour each code; compare as text so 2 typed as a number still matches.
    top_left = f"{day_col(1)}{FIRST_DATA_ROW}"
    for code, _, _, colour in CODES:
        ws.conditional_formatting.add(
            grid,
            FormulaRule(
                formula=[f'TRIM(UPPER({top_left}&""))="{code}"'],
                fill=PatternFill("solid", fgColor=colour),
                font=Font(bold=True),
            ),
        )
    # Grey out days the month doesn't have, and tint weekends in the day rows.
    head = f"{day_col(1)}3:{day_col(DAYS)}{HEADER_ROW}"
    first = f"{day_col(1)}$4"
    in_month = f"MONTH(DATE(YEAR($B$2),MONTH($B$2),{first}))=MONTH($B$2)"
    ws.conditional_formatting.add(
        head, FormulaRule(formula=[f"NOT({in_month})"], fill=PatternFill("solid", fgColor="7F7F7F"))
    )
    ws.conditional_formatting.add(
        head,
        FormulaRule(
            formula=[f"AND({in_month},WEEKDAY(DATE(YEAR($B$2),MONTH($B$2),{first}),2)>5)"],
            fill=PatternFill("solid", fgColor="C00000"),
            font=Font(color=WHITE, bold=True),
        ),
    )

    ws.freeze_panes = ws.cell(FIRST_DATA_ROW, FIRST_DAY_COL)
    ws.sheet_view.zoomScale = 100
    # Print the whole month across one landscape page width.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"{HEADER_ROW - 1}:{HEADER_ROW}"


def build_help_sheet(ws):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 60

    ws["A1"] = "วิธีใช้ไฟล์ต้นแบบตารางกะ"
    ws["A1"].font = font(bold=True, size=14, color=NAVY)

    steps = [
        "1. ไปที่แท็บ “Shift Planning” แล้วใส่วันที่ 1 ของเดือนในช่อง B2 (เช่น 1/11/2026)",
        "2. กรอก 1 คนต่อ 1 แถว: Employee ID, Name (ชื่อเล่นที่ใช้ในทีม), Title, Squad",
        "3. ใส่รหัสกะในช่องของแต่ละวัน (เลือกจากดรอปดาวน์ได้) · เว้นว่าง = ยังไม่ลงกะ",
        "4. ห้ามแก้แถวหัวตาราง (แถว 4) และห้ามเปลี่ยนเลขวันที่ 1–31 เป็นข้อความอื่น",
        "5. นำเข้า: ใน LiveOps Scheduler ไปที่ Staff Shift → “นำเข้ากะจาก Google Sheet” → อัปโหลดไฟล์ .xlsx",
        "   (ถ้าใช้บน Google Sheets: File → Download → Microsoft Excel (.xlsx) แล้วอัปโหลด)",
        "6. ตรวจหน้าตัวอย่างก่อนกดยืนยัน — การยืนยันจะแทนที่กะทั้งเดือนของคนที่อยู่ในไฟล์",
    ]
    for i, s in enumerate(steps, start=3):
        ws.cell(i, 1, s).font = font(size=11)

    r = len(steps) + 5
    ws.cell(r, 1, "รหัสกะ").font = font(bold=True, size=12, color=NAVY)
    r += 1
    for i, h in enumerate(["รหัส", "ความหมาย", "เวลา"], start=1):
        c = ws.cell(r, i, h)
        c.font = font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = BOX
    for code, meaning, hours, colour in CODES:
        r += 1
        ws.cell(r, 1, code).fill = PatternFill("solid", fgColor=colour)
        ws.cell(r, 1).font = font(bold=True)
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, meaning)
        ws.cell(r, 3, hours)
        for c in range(1, 4):
            ws.cell(r, c).border = BOX
    r += 1
    ws.cell(r, 1, "ว่าง").alignment = Alignment(horizontal="center")
    ws.cell(r, 2, "ยังไม่ลงกะ (ระบบจะข้ามช่องนี้)")
    for c in range(1, 4):
        ws.cell(r, c).border = BOX

    r += 2
    ws.cell(r, 1, "การจับคู่พนักงาน").font = font(bold=True, size=12, color=NAVY)
    notes = [
        "• ระบบจับคู่ด้วย Employee ID ก่อน — ชื่อเล่นเปลี่ยนก็ยังจับคู่ถูกคน",
        "• คนที่ยังไม่มีในระบบจะถูกเพิ่มเป็นพนักงานใหม่ พร้อมตำแหน่งและ Squad จากไฟล์",
        "• แท็บ “ตัวอย่าง” มีไว้ดูเท่านั้น ระบบจะอ่านแท็บ “Shift Planning”",
    ]
    for i, s in enumerate(notes, start=r + 1):
        ws.cell(i, 1, s)


def main():
    month = date(2026, 11, 1)
    wb = Workbook()
    shift = wb.active
    shift.title = "Shift Planning"
    build_shift_sheet(shift, month)

    build_help_sheet(wb.create_sheet("วิธีใช้"))

    pattern = ["2", "2", "4", "3", "1", "1", "2"]
    example = [
        ("110112", "M", "Producer Team Lead", "", ["1", "OF", "OF", "OF", "OF", "OF", "1"] * 4 + ["1", "OF"]),
        ("110215", "First", "Producer", "Uniqlo", [pattern[(d + 1) % 7] for d in range(30)]),
        ("110319", "Bookbik", "Jr. Producer", "Uniqlo", [pattern[(d + 3) % 7] for d in range(30)]),
        ("110568", "Donut", "Jr. Producer", "Adidas Ch1 (TikTok)", ["VL", "VL"] + [pattern[(d + 5) % 7] for d in range(28)]),
        ("110640", "Arm", "Jr. Producer", "Adidas Ch1 (TikTok)", [pattern[d % 7] for d in range(20)] + [""] * 10),
    ]
    sample = wb.create_sheet("ตัวอย่าง")
    build_shift_sheet(sample, month, example)
    sample["A1"] = "ตัวอย่าง — ข้อมูลสมมติ (ระบบไม่นำเข้าแท็บนี้)"

    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
